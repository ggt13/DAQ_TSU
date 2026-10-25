"""Generador común de la hoja de resultados (se sube a Drive y se abre como hoja de cálculo de Google).

Pestañas:
- Resultados: nombre y código que escribe el alumno; verificación contra el CSV; resultados; predicho contra
  obtenido; preguntas del reporte; tabla para la gráfica (encabezados con el nombre, así la leyenda lo muestra).
- CSV: aquí se importa el archivo que descarga el cuaderno (Archivo → Importar → Reemplazar la hoja actual).

Cada dato se busca por su etiqueta en la pestaña CSV, no por su renglón. Las celdas vacías de la tabla dan #N/A
para que la gráfica no dibuje ceros.
"""
import csv

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL = Font(color="0000FF")
NEGRITA = Font(bold=True)
GRIS = Font(italic=True, color="666666")
CAPTURA = PatternFill("solid", fgColor="FFF2CC")
ROJO = PatternFill("solid", fgColor="F4CCCC")
VERDE = PatternFill("solid", fgColor="D9EAD3")


def busca(etiqueta):
    return f'=IFERROR(VLOOKUP("{etiqueta}",CSV!$A:$B,2,FALSE),"")'


def _num(x):
    try:
        return float(x)
    except ValueError:
        return x


def generar(salida, resultados, comparacion, preguntas, tabla, grafica, renglones,
            metadatos=(("Fuente de los datos", "fuente"),), csv_ejemplo=None):
    """
    resultados: [(etiqueta_en_csv, texto, formato)].
    comparacion: (etiqueta_predicho, etiqueta_obtenido, texto) o None.
    preguntas: [texto].
    tabla: {"encabezado_csv": nombre de la primera columna de la tabla en el CSV,
            "x": (columna_csv, título), "series": [(columna_csv, prefijo_de_leyenda)]}.
    grafica: {"titulo", "x", "y", "tipo": "linea" | "dispersion"}.
    """
    wb = Workbook()
    R = wb.active
    R.title = "Resultados"
    C = wb.create_sheet("CSV")
    if csv_ejemplo:
        for f in csv.reader(open(csv_ejemplo, encoding="utf-8")):
            C.append([_num(x) for x in f])
    else:
        C["A1"] = "Importa aquí tu CSV: Archivo → Importar → Subir → Reemplazar la hoja actual"

    R.column_dimensions["A"].width = 48
    R.column_dimensions["B"].width = 28
    R["A1"], R["A2"] = "Nombre completo ✎", "Código de alumno ✎"
    for c in ("B1", "B2"):
        R[c].font, R[c].fill = AZUL, CAPTURA
    filas = [("Código en el CSV", busca("codigo")),
             ("Verificación", '=IF(B3="","Importa tu CSV",IF(B2="","Escribe tu código",'
                              'IF(TRIM(B3&"")=TRIM(B2&""),"OK","NO COINCIDE")))'),
             ("Huella del cuaderno", busca("huella")),
             ("Sesión", busca("sesion"))] + [(t, busca(k)) for t, k in metadatos]
    f = 3
    for a, b in filas:
        R.cell(f, 1, a)
        R.cell(f, 2, b)
        f += 1
    R.conditional_formatting.add("B4", FormulaRule(formula=['B4="NO COINCIDE"'], fill=ROJO))
    R.conditional_formatting.add("B4", FormulaRule(formula=['B4="OK"'], fill=VERDE))
    R.conditional_formatting.add("B7", FormulaRule(formula=['ISNUMBER(SEARCH("SINT",B7))'], fill=ROJO))

    f += 1
    R.cell(f, 1, "Tus resultados").font = NEGRITA
    fila_de = {}
    for clave, texto, formato in resultados:
        f += 1
        R.cell(f, 1, texto)
        R.cell(f, 2, busca("resultado:" + clave)).number_format = formato
        fila_de[clave] = f
    if comparacion:
        kp, ko, texto = comparacion
        p, o = f"B{fila_de[kp]}", f"B{fila_de[ko]}"
        f += 2
        R.cell(f, 1, "Predicho contra obtenido").font = NEGRITA
        f += 1
        R.cell(f, 1, texto)
        R.cell(f, 2, f'=IF(OR({p}="",{o}=""),"",({o}-{p})/{p}*100)').number_format = "0.0"

    f += 2
    R.cell(f, 1, "Preguntas del reporte ✎").font = NEGRITA
    for p in preguntas:
        f += 1
        R.cell(f, 1, p).alignment = Alignment(wrap_text=True, vertical="top")
        R.cell(f, 2).fill = CAPTURA
        R.cell(f, 2).alignment = Alignment(wrap_text=True, vertical="top")
        R.merge_cells(start_row=f, start_column=2, end_row=f, end_column=7)
        R.row_dimensions[f].height = 75
    f += 2
    R.cell(f, 1, "La gráfica se arma sola al importar tu CSV. Cambia su título por tu nombre (doble clic en el "
                 "título); la leyenda ya lo toma de B1.").font = GRIS
    fila_grafica = f + 1

    # Tabla para la gráfica, a partir de la columna I. H2 guarda el renglón del encabezado en la pestaña CSV.
    R["H1"], R["H2"] = "Encabezado en CSV", f'=IFERROR(MATCH("{tabla["encabezado_csv"]}",CSV!$A:$A,0),"")'
    R["H1"].font = GRIS
    cols = [tabla["x"]] + tabla["series"]
    for j, (col_csv, titulo) in enumerate(cols):
        letra = get_column_letter(9 + j)
        R.column_dimensions[letra].width = 24
        R[f"{letra}1"] = titulo if j == 0 else f'="{titulo} · "&$B$1'
        R[f"{letra}1"].font = NEGRITA
        for k in range(1, renglones + 1):
            ref = f"INDEX(CSV!${col_csv}:${col_csv},$H$2+{k})"
            R[f"{letra}{1 + k}"] = f'=IF($H$2="",NA(),IF({ref}="",NA(),{ref}*1))'
    ultimo = 1 + renglones
    fin = get_column_letter(9 + len(cols) - 1)
    R.conditional_formatting.add(f"I2:{fin}{ultimo}", FormulaRule(formula=["ISNA(I2)"], font=Font(color="FFFFFF")))

    if grafica.get("tipo") == "dispersion":
        g = ScatterChart()
        xs = Reference(R, min_col=9, min_row=2, max_row=ultimo)
        for j in range(1, len(cols)):
            s = Series(Reference(R, min_col=9 + j, min_row=1, max_row=ultimo), xs, title_from_data=True)
            s.marker.symbol = "circle"
            g.series.append(s)
    else:
        g = LineChart()
        g.add_data(Reference(R, min_col=10, max_col=9 + len(cols) - 1, min_row=1, max_row=ultimo), titles_from_data=True)
        g.set_categories(Reference(R, min_col=9, min_row=2, max_row=ultimo))
    g.title = grafica["titulo"] + " · (tu nombre)"
    g.x_axis.title = grafica["x"]
    g.y_axis.title = grafica["y"]
    g.height, g.width = 9, 22
    R.add_chart(g, f"A{fila_grafica}")
    wb.save(salida)
    return salida
