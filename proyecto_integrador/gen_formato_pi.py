# -*- coding: utf-8 -*-
"""Genera el formato del alumno del proyecto integrador: F_Proyecto_integrador_DAQ_2027_GGT.docx

Todo el contenido sale de datos_pi.py (fuente única); aquí sólo está la maqueta, con el aspecto del formato del
proyecto de Sensores 2027: encabezado, Nombre/Grupo/Fecha, recuadro «Qué es este documento», dos sesiones por página
(campos a la izquierda, recuadro de la gráfica a la derecha) y cuatro secciones finales.

Uso:  python gen_formato_pi.py            -> F_Proyecto_integrador_DAQ_2027_GGT.docx junto a este archivo
      python gen_formato_pi.py --pdf      -> además el PDF (LibreOffice)
"""
import os
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import datos_pi as D  # noqa: E402

AZUL, AMBAR, GRIS, TEXTO = "00567C", "C9741F", "4A5560", "1A1A1A"
CLARO, CREMA, LINEA = "EEF4F7", "FDF3E7", "9AA7B1"
FUENTE = "Arial"
ANCHO = 17.9            # cm útiles (carta, márgenes de 1.85 cm)


# ----------------------------------------------------------------------------------------------- utilidades

ORDEN = {
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark"],
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
            "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
            "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
              "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"],
}


def inserta(padre, hijo):
    """Inserta `hijo` en `padre` en el orden que exige el esquema de Word (si no, Word puede marcar el archivo)."""
    orden = ORDEN[padre.tag.split("}")[1]]
    k = orden.index(hijo.tag.split("}")[1])
    for i, e in enumerate(list(padre)):
        nombre = e.tag.split("}")[1]
        if nombre in orden and orden.index(nombre) > k:
            padre.insert(i, hijo)
            return hijo
    padre.append(hijo)
    return hijo

def rgb(h):
    return RGBColor.from_string(h)


def texto(p, t, pt=10.5, bold=False, italic=False, color=TEXTO, mono=False):
    r = p.add_run(t)
    r.font.size = Pt(pt); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color)
    r.font.name = "Liberation Mono" if mono else FUENTE
    r._element.rPr.rFonts.set(qn("w:eastAsia"), r.font.name)
    return r


def ecuacion(p, ec, pt=10.5):
    """Escribe una ecuación de datos_pi.py con los subíndices «_x» como subíndices reales y sin cortes de renglón."""
    import re
    for k, trozo in enumerate(re.split(r"_([A-Za-z0-9]+)", ec)):
        if not trozo:
            continue
        r = texto(p, trozo.replace(" ", "\u00a0"), pt - 2 if k % 2 else pt, bold=True)
        r.font.subscript = bool(k % 2)


def parrafo(destino, t="", pt=10.5, bold=False, italic=False, color=TEXTO, align=None, antes=0, despues=4,
            mono=False, keep=False):
    p = destino.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(antes); pf.space_after = Pt(despues); pf.line_spacing = 1.08
    if keep:
        pf.keep_with_next = True
    if align:
        p.alignment = align
    if t:
        texto(p, t, pt, bold, italic, color, mono)
    return p


def primer_parrafo(celda):
    """El párrafo vacío que trae cada celda; se reutiliza para no dejar renglones de más."""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(0)
    return p


def sombra(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    inserta(tcPr, shd)


def bordes_celda(celda, estilo="single", color=LINEA, grosor=6, lados=("top", "left", "bottom", "right")):
    tcPr = celda._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = inserta(tcPr, OxmlElement("w:tcBorders"))
    for lado in ("top", "left", "bottom", "right"):
        e = b.find(qn(f"w:{lado}"))
        if e is None:
            e = OxmlElement(f"w:{lado}"); b.append(e)
        if lado in lados:
            e.set(qn("w:val"), estilo); e.set(qn("w:sz"), str(grosor)); e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")


def margenes_celda(celda, cm=0.18):
    tcPr = celda._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}"); e.set(qn("w:w"), str(int(cm * 567))); e.set(qn("w:type"), "dxa")
        m.append(e)
    vieja = tcPr.find(qn("w:tcMar"))
    if vieja is not None:
        tcPr.remove(vieja)
    inserta(tcPr, m)


def tabla(destino, filas, anchos, bordes=True):
    t = destino.add_table(rows=filas, cols=len(anchos))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for fila in t.rows:
        for c, w in zip(fila.cells, anchos):
            c.width = Cm(w)
            margenes_celda(c)
            if bordes:
                bordes_celda(c)
            else:
                bordes_celda(c, lados=())
    grid = t._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), anchos):
        col.set(qn("w:w"), str(int(w * 567)))
    return t


def alto(fila, cm, exacto=False):
    fila.height = Cm(cm)
    fila.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exacto else WD_ROW_HEIGHT_RULE.AT_LEAST


def no_partir(fila):
    trPr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit"); trPr.append(e)


def repetir_encabezado(fila):
    trPr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader"); trPr.append(e)


def salto(doc):
    """Salto de página con un párrafo mínimo (no ocupa un renglón al inicio de la página nueva)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.page_break_before = True; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = Pt(1)
    texto(p, "", 1)


def titulo_seccion(doc, t, sub=None):
    p = parrafo(doc, t, pt=16, bold=True, color=AZUL, despues=2, keep=True)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12"); bot.set(qn("w:color"), AMBAR); bot.set(qn("w:space"), "2")
    bdr.append(bot); inserta(pPr, bdr)
    if sub:
        parrafo(doc, sub, pt=10.5, color=GRIS, antes=4, despues=8, keep=True)


def recuadro(doc, titulo, alto_cm, nota=None, ancho=ANCHO):
    """Un recuadro punteado para pegar una gráfica o escribir; el título va dentro, arriba."""
    t = tabla(doc, 1, [ancho], bordes=False)
    c = t.rows[0].cells[0]
    bordes_celda(c, estilo="dashed", color=GRIS, grosor=8)
    alto(t.rows[0], alto_cm); no_partir(t.rows[0])
    p = primer_parrafo(c)
    texto(p, titulo, 10.5, bold=True, color=AZUL)
    if nota:
        texto(primer_parrafo_extra(c), nota, 10, color=GRIS)
    return t


def primer_parrafo_extra(celda):
    p = celda.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return p


# ----------------------------------------------------------------------------------------------- documento
def pie(seccion):
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texto(p, f"Proyecto integrador · {D.CURSO} · 2027 · página ", 9, color=GRIS)
    r = p.add_run()
    for tipo, val in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tipo); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = val; r._r.append(e)
    r.font.size = Pt(9); r.font.color.rgb = rgb(GRIS); r.font.name = FUENTE


def portada(doc):
    parrafo(doc, "Proyecto integrador", pt=24, bold=True, color=AZUL, despues=0)
    parrafo(doc, D.CURSO, pt=15, color=TEXTO, despues=0)
    parrafo(doc, D.PROGRAMA, pt=10.5, color=GRIS, despues=10)

    t = tabla(doc, 2, [9.0, 4.4, 4.5], bordes=False)
    for i, lab in enumerate(("Nombre", "Código · Grupo", "Fecha")):
        c0, c1 = t.rows[0].cells[i], t.rows[1].cells[i]
        bordes_celda(c0, color=GRIS, grosor=8, lados=("bottom",))
        bordes_celda(c1, lados=())
        p = primer_parrafo(c1); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        texto(p, lab, 9.5, color=GRIS)
    alto(t.rows[0], 0.95)
    parrafo(doc, despues=6)

    # Qué es este documento
    t = tabla(doc, 1, [ANCHO], bordes=False)
    c = t.rows[0].cells[0]; sombra(c, CLARO); bordes_celda(c, color=AZUL, grosor=8, lados=("left",))
    margenes_celda(c, 0.3)
    p = primer_parrafo(c); texto(p, "Qué es este documento", 12, bold=True, color=AZUL)
    p = primer_parrafo_extra(c)
    texto(p, "Es tu proyecto integrador y es individual. Al cierre de cada sesión llenas cuatro campos; al final "
             "completas cuatro secciones: tu plan de adquisición para una variable de tu área, tu registro real "
             "contra ese plan, la autoevaluación y la reflexión. Cada semana se revisa en los 10 minutos de revisión, "
             "pero sólo se califica el reporte final.", 10.5)
    p = primer_parrafo_extra(c); p.paragraph_format.space_before = Pt(4)
    texto(p, "Los cuatro campos de cada sesión", 10.5, bold=True, color=AZUL)
    for nombre, que in D.CAMPOS_SESION:
        p = primer_parrafo_extra(c); p.paragraph_format.left_indent = Cm(0.4)
        texto(p, nombre + ". ", 10.5, bold=True); texto(p, que, 10.5)
    p = primer_parrafo_extra(c); p.paragraph_format.space_before = Pt(4)
    texto(p, "Sin datos confidenciales. ", 10.5, bold=True, color=AMBAR)
    texto(p, "No anotes marcas, números de parte ni valores de la línea. Si tu área no tiene relación con el tema de "
             "una sesión, dilo y explica por qué; esa respuesta cuenta igual.", 10.5)
    parrafo(doc, despues=4)

    # Cómo se califica
    parrafo(doc, "Cómo se califica · 30 % de la calificación del curso", pt=12, bold=True, color=AZUL, despues=3,
            keep=True)
    t = tabla(doc, 1 + len(D.RUBRICA), [4.6, 1.6, 11.7])
    for i, h in enumerate(("Componente", "Peso", "Qué se revisa")):
        c = t.rows[0].cells[i]; sombra(c, AZUL)
        texto(primer_parrafo(c), h, 10, bold=True, color="FFFFFF")
    for k, (comp, peso, que) in enumerate(D.RUBRICA, 1):
        cs = t.rows[k].cells
        texto(primer_parrafo(cs[0]), comp, 10.5, bold=True)
        p = primer_parrafo(cs[1]); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; texto(p, f"{peso} %", 10.5, bold=True,
                                                                                   color=AMBAR)
        texto(primer_parrafo(cs[2]), que, 10)
    parrafo(doc, despues=4)
    parrafo(doc, D.ENTREGA, pt=10, color=GRIS)


def sesion(doc, s):
    # banda del tema
    t = tabla(doc, 1, [1.3, ANCHO - 1.3], bordes=False)
    c0, c1 = t.rows[0].cells
    sombra(c0, AZUL); sombra(c1, CLARO)
    c0.vertical_alignment = c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = primer_parrafo(c0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texto(p, str(s["n"]), 18, bold=True, color="FFFFFF")
    p = primer_parrafo(c1)
    texto(p, s["tema"], 12.5, bold=True, color=AZUL); texto(p, "  ·  " + s["que"], 10.5, color=GRIS)
    alto(t.rows[0], 0.95)
    no_partir(t.rows[0])
    for p in c1.paragraphs + c0.paragraphs:
        p.paragraph_format.keep_with_next = True

    (n1, h1), (n2, h2), (n3, h3), (n4, h4) = D.CAMPOS_SESION
    t = tabla(doc, 1, [11.4, ANCHO - 11.4], bordes=False)
    no_partir(t.rows[0]); alto(t.rows[0], 8.4)
    iz, de = t.rows[0].cells
    bordes_celda(iz, color=LINEA, lados=("left", "bottom", "right"))
    bordes_celda(de, estilo="dashed", color=GRIS, grosor=8)

    def campo(celda, nombre, pista, primero=False):
        p = primer_parrafo(celda) if primero else primer_parrafo_extra(celda)
        p.paragraph_format.space_before = Pt(0 if primero else 6)
        texto(p, nombre, 10.5, bold=True, color=AZUL)
        p = primer_parrafo_extra(celda); texto(p, pista, 10, italic=True, color=GRIS)

    def renglones(celda, k):
        """Renglones punteados para escribir (tabulador con guía de puntos; los bordes de párrafo se fusionan)."""
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        for _ in range(k):
            p = primer_parrafo_extra(celda)
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(10.9), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS)
            texto(p, "\t\u200b", 10, color=LINEA)

    campo(iz, n1, h1, primero=True); renglones(iz, 2)
    p = primer_parrafo_extra(iz); p.paragraph_format.space_before = Pt(6)
    texto(p, n2, 10.5, bold=True, color=AZUL); texto(p, "  " + h2, 10, italic=True, color=GRIS)
    ecs = [f.strip() for f in s["ecuacion"].split("   ") if f.strip()]
    for k in range(0, len(ecs), 2):                      # dos ecuaciones por renglón
        p = primer_parrafo_extra(iz); p.paragraph_format.space_before = Pt(3 if k == 0 else 0)
        p.paragraph_format.left_indent = Cm(0.3)
        for j, ec in enumerate(ecs[k:k + 2]):
            if j:
                texto(p, "     ", 10.5)
            ecuacion(p, ec)
    p = primer_parrafo_extra(iz); p.paragraph_format.space_before = Pt(4)
    texto(p, "Mi sustitución y mi resultado:", 10.5, bold=True)
    renglones(iz, 2)
    p = primer_parrafo_extra(iz); p.paragraph_format.space_before = Pt(6)
    texto(p, n3, 10.5, bold=True, color=AZUL)
    p = primer_parrafo_extra(iz); texto(p, s["pregunta"], 10, italic=True, color=GRIS)
    renglones(iz, 2)

    p = primer_parrafo(de); texto(p, n4, 10.5, bold=True, color=AZUL)
    p = primer_parrafo_extra(de); texto(p, s["grafica"] + ".", 10, italic=True, color=GRIS)
    p = primer_parrafo_extra(de); texto(p, h4, 10, color=GRIS)


def ocho_sesiones(doc):
    for k, s in enumerate(D.SESIONES):
        if k % 2 == 0:
            salto(doc)       # sin título de sección: dos sesiones llenan la página; la banda numerada las nombra
        sesion(doc, s)
        if k % 2 == 0:
            parrafo(doc, despues=0).paragraph_format.line_spacing = Pt(10)


def tabla_doble(doc, filas, col_que, col_mia):
    """Qué · Unidad · En el cuaderno · lo del alumno para su variable · lo de su registro, si es de otra variable."""
    anchos = [5.9, 2.4, 3.4, 3.1, 3.1]
    t = tabla(doc, 1 + len(filas), anchos)
    for i, h in enumerate((col_que, "Unidad", "En el cuaderno", col_mia, "Para mi registro, si es de otra variable")):
        c = t.rows[0].cells[i]; sombra(c, AZUL)
        texto(primer_parrafo(c), h, 10, bold=True, color="FFFFFF")
    repetir_encabezado(t.rows[0]); no_partir(t.rows[0])
    for k, (var, que, uni) in enumerate(filas, 1):
        fila = t.rows[k]; alto(fila, 0.95); no_partir(fila)
        cs = fila.cells
        texto(primer_parrafo(cs[0]), que, 10.5)
        texto(primer_parrafo(cs[1]), uni, 10, color=GRIS)
        texto(primer_parrafo(cs[2]), var, 10, color=GRIS, mono=True)
        sombra(cs[4], "F6F8FA")
    return t


def plan(doc):
    salto(doc)
    titulo_seccion(doc, "1 · Mi plan de adquisición",
                   "El plan se completa al final del curso, para una variable de tu área: en cada campo escribe lo "
                   "que propones; la columna «Sesión» dice de dónde lo tomas. Las predicciones son la excepción: se "
                   "escriben con fecha en la semana 2, antes de que empiece tu registro (si usas una campaña, antes "
                   "de abrir su archivo).")
    anchos = [3.3, 5.0, 1.6, ANCHO - 9.9]
    t = tabla(doc, 1 + len(D.PLAN), anchos)
    for i, h in enumerate(("Campo", "Qué escribes", "Sesión", "Mi plan")):
        c = t.rows[0].cells[i]; sombra(c, AZUL)
        texto(primer_parrafo(c), h, 10, bold=True, color="FFFFFF")
    repetir_encabezado(t.rows[0])
    for k, (campo, que, ses) in enumerate(D.PLAN, 1):
        fila = t.rows[k]; alto(fila, 2.3); no_partir(fila)
        cs = fila.cells
        sombra(cs[0], CLARO)
        texto(primer_parrafo(cs[0]), campo, 10.5, bold=True, color=AZUL)
        texto(primer_parrafo(cs[1]), que, 10, color=GRIS)
        p = primer_parrafo(cs[2]); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; texto(p, ses, 10, bold=True, color=AMBAR)

    # Los números del plan
    parrafo(doc, "Los números de mi plan (los copias al cuaderno del proyecto)", pt=12.5, bold=True, color=AZUL,
            antes=12, despues=6, keep=True)
    tabla_doble(doc, [(var, que, uni) for var, que, uni, _ in D.NUMEROS_PLAN], "Qué", "Para mi variable")
    parrafo(doc, D.OTRA_VARIABLE, pt=9, italic=True, color=GRIS, antes=4, despues=0)

    # La cadena
    parrafo(doc, "La cadena de mi plan", pt=12.5, bold=True, color=AZUL, antes=12, despues=2, keep=True)
    parrafo(doc, "Escribe en cada caja lo que propones. Si una queda vacía, es la parte que aún no sabes resolver.",
            pt=10, color=GRIS, despues=6, keep=True)
    n = len(D.CADENA); flecha = 0.6
    largos = [len(nombre) for nombre, _ in D.CADENA]
    ancha = max(range(n), key=lambda i: largos[i])          # la caja del nombre más largo va más ancha
    libre = ANCHO - flecha * (n - 1)
    w_ancha = 3.4
    w_normal = (libre - w_ancha) / (n - 1)
    anchos = []
    for i in range(n):
        anchos.append(w_ancha if i == ancha else w_normal)
        if i < n - 1:
            anchos.append(flecha)
    t = tabla(doc, 1, anchos, bordes=False)
    alto(t.rows[0], 4.2); no_partir(t.rows[0])
    for i, (nombre, ses) in enumerate(D.CADENA):
        c = t.rows[0].cells[2 * i]
        bordes_celda(c, color=AZUL, grosor=10)
        p = primer_parrafo(c); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        texto(p, nombre, 10, bold=True, color=AZUL)
        p = primer_parrafo_extra(c); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        texto(p, ses, 9.5, color=AMBAR)
        if i < n - 1:
            f = t.rows[0].cells[2 * i + 1]
            f.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margenes_celda(f, 0.0)
            p = primer_parrafo(f); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            texto(p, "→", 16, bold=True, color=AMBAR)

    # Predicciones
    parrafo(doc, "Mis predicciones, en la semana 2, antes del registro", pt=12.5, bold=True, color=AZUL, antes=14,
            despues=2, keep=True)
    parrafo(doc, "Escríbelas con fecha en la semana 2, antes de que empiece tu registro (si usas una campaña, antes de "
                 "abrir su archivo). El cuaderno del proyecto las compara con lo obtenido.",
            pt=10, color=GRIS, despues=6, keep=True)
    tabla_doble(doc, D.PREDICCIONES, "Qué predices", "Mi predicción")
    p = parrafo(doc, antes=8, despues=0)
    texto(p, "Fecha en que escribí mis predicciones: ", 10.5, bold=True); texto(p, "_" * 30, 10.5, color=LINEA)


def registro(doc):
    salto(doc)
    titulo_seccion(doc, "2 · Mi registro contra mi plan",
                   "Un registro real analizado con el cuaderno del proyecto. Elige una fuente, llena la ficha y pega "
                   "lo que produce el cuaderno.")
    parrafo(doc, "La fuente de mi registro", pt=12.5, bold=True, color=AZUL, despues=4, keep=True)
    t = tabla(doc, len(D.FUENTES), [1.0, ANCHO - 1.0], bordes=False)
    for k, (nombre, desc) in enumerate(D.FUENTES):
        c0, c1 = t.rows[k].cells
        p = primer_parrafo(c0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; texto(p, "☐", 15, color=AZUL)
        p = primer_parrafo(c1); texto(p, nombre + ". ", 10.5, bold=True); texto(p, desc, 10, color=GRIS)
        p.paragraph_format.space_after = Pt(5)

    parrafo(doc, "Ficha del registro", pt=12.5, bold=True, color=AZUL, antes=10, despues=4, keep=True)
    t = tabla(doc, len(D.FICHA_REGISTRO), [5.2, ANCHO - 5.2])
    for k, campo in enumerate(D.FICHA_REGISTRO):
        fila = t.rows[k]; no_partir(fila)
        alto(fila, 2.6 if k == len(D.FICHA_REGISTRO) - 1 else 1.0)
        sombra(fila.cells[0], CLARO)
        texto(primer_parrafo(fila.cells[0]), campo, 10.5, bold=True, color=AZUL)

    parrafo(doc, "Lo que produce el cuaderno del proyecto", pt=12.5, bold=True, color=AZUL, antes=12, despues=2,
            keep=True)
    parrafo(doc, "Pega las gráficas de la última celda del cuaderno, las que llevan tu huella, y la tabla.", pt=10,
            color=GRIS, despues=6, keep=True)
    for tit, nota in (("Gráfica 1 · Mi registro y sus huecos", "Duración, intervalo efectivo y huecos."),
                      ("Gráfica 2 · Filtro y tendencia", "El registro, el promedio móvil de tu ventana y la recta."),
                      ("Gráfica 3 · Alarma con confirmación",
                       "Tu límite, el tiempo de confirmación, la histéresis y las alarmas.")):
        recuadro(doc, tit, 7.2, nota)
        parrafo(doc, despues=4)
    recuadro(doc, "Tabla · Plan contra registro", 6.0, "Lo que predijiste contra lo que obtuviste, con la diferencia.")
    parrafo(doc, despues=4)
    recuadro(doc, "Conclusión y decisión", 7.0,
             "¿Coincidió tu registro con tu plan? Explica las diferencias con tus números y escribe la decisión que "
             "tomarías y con qué tipo de mantenimiento: correctivo, preventivo, basado en condición o predictivo.")
    p = parrafo(doc, antes=10, despues=0)
    texto(p, "Huella del cuaderno del proyecto: ", 11, bold=True, color=AZUL)
    texto(p, "_" * 16, 11, color=LINEA)


def autoevaluacion(doc):
    salto(doc)
    titulo_seccion(doc, "3 · Autoevaluación",
                   "Marca con una X dónde estás hoy. Se califica que la tabla esté completa y sea honesta, no el "
                   "nivel que marques.")
    niv = 2.3
    t = tabla(doc, 1 + len(D.CAPACIDADES), [ANCHO - 4 * niv] + [niv] * 4)
    c = t.rows[0].cells[0]; sombra(c, AZUL); texto(primer_parrafo(c), "Al terminar el curso puedo…", 10, bold=True,
                                                     color="FFFFFF")
    for i, n in enumerate(D.NIVELES, 1):
        c = t.rows[0].cells[i]; sombra(c, AZUL)
        p = primer_parrafo(c); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; texto(p, n, 9.5, bold=True, color="FFFFFF")
    for k, cap in enumerate(D.CAPACIDADES, 1):
        fila = t.rows[k]; alto(fila, 1.05); no_partir(fila)
        texto(primer_parrafo(fila.cells[0]), f"{k}. {cap}", 10.5)
        if k % 2 == 0:
            for c in fila.cells:
                sombra(c, "F6F8FA")

    salto(doc)
    titulo_seccion(doc, "4 · Reflexión del curso", "Respuestas concretas, con ejemplos de tu trabajo o de tu registro.")
    for q in D.REFLEXION:
        parrafo(doc, q, pt=10.5, bold=True, color=AZUL, antes=6, despues=3, keep=True)
        t = tabla(doc, 1, [ANCHO])
        alto(t.rows[0], 3.7); no_partir(t.rows[0])
    p = parrafo(doc, antes=14, despues=0)
    texto(p, D.ENTREGA, 10, italic=True, color=GRIS)


def construye(ruta):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FUENTE; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(1.85)
    sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.5)
    sec.footer_distance = Cm(0.8)
    pie(sec)
    doc.core_properties.title = f"Proyecto integrador · {D.CURSO} · 2027"
    doc.core_properties.author = "Guillermo García Torales"

    portada(doc)
    ocho_sesiones(doc)
    plan(doc)
    registro(doc)
    autoevaluacion(doc)
    z = doc.settings.element.find(qn("w:zoom"))
    if z is not None and z.get(qn("w:percent")) is None:
        z.set(qn("w:percent"), "100")
    doc.save(ruta)
    return ruta


if __name__ == "__main__":
    salida = construye(os.path.join(AQUI, D.FORMATO))
    print(salida)
    if "--pdf" in sys.argv:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", AQUI, salida], check=True,
                       capture_output=True, timeout=180)
        print(salida.replace(".docx", ".pdf"))
