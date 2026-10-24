"""Genera la hoja de resultados de muestra (sesión 0) a partir de un CSV firmado.

Uso:
    python gen_hoja_plantilla.py CSV_DE_EJEMPLO   -> S0_AC_PRACT_Hoja_de_resultados_PLANTILLA.xlsx

Al subirla a Drive se abre como hoja de cálculo de Google. Pestañas:
- Resultados: nombre y código que escribe el alumno, verificación contra el CSV, predicho contra obtenido,
  tabla para la gráfica (los encabezados llevan el nombre, así la leyenda lo muestra) y las preguntas del reporte.
- CSV: aquí se importa el archivo que descarga el cuaderno (Archivo → Importar → Reemplazar la hoja actual).
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, "S0_AC_PRACT_Hoja_de_resultados_PLANTILLA.xlsx")
AZUL = Font(color="0000FF")                       # lo que escribe el alumno
NEGRITA = Font(bold=True)
CAPTURA = PatternFill("solid", fgColor="FFF2CC")   # celdas a llenar
ROJO = PatternFill("solid", fgColor="F4CCCC")
VERDE = PatternFill("solid", fgColor="D9EAD3")

ETIQUETAS = {"lsb_C": "LSB del convertidor (°C)",
             "error_predicho_C": "Error máximo predicho, ejercicio 2 (°C)",
             "error_obtenido_C": "Error máximo obtenido con datos reales (°C)"}


def num(x):
    try:
        return float(x)
    except ValueError:
        return x


def generar(ruta_csv):
    filas = list(csv.reader(open(ruta_csv, encoding="utf-8")))
    n_meta = next(i for i, f in enumerate(filas) if not f)        # renglón vacío que separa la tabla
    fila_enc = n_meta + 2                                          # 1-indexado en la hoja
    datos = filas[n_meta + 2:]
    res = [(i + 1, f[0].split(":", 1)[1]) for i, f in enumerate(filas[:n_meta]) if f[0].startswith("resultado:")]

    wb = Workbook()
    R = wb.active
    R.title = "Resultados"
    C = wb.create_sheet("CSV")
    for f in filas:
        C.append([num(x) for x in f])

    R.column_dimensions["A"].width = 44
    R.column_dimensions["B"].width = 26
    for col in "IJK":
        R.column_dimensions[col].width = 22
    R["A1"], R["A2"] = "Nombre completo", "Código de alumno"
    for c in ("B1", "B2"):
        R[c].font, R[c].fill = AZUL, CAPTURA
    R["A3"], R["B3"] = "Código en el CSV", "=CSV!B2"
    R["A4"] = "Verificación"
    R["B4"] = '=IF(B2="","Escribe tu código",IF(TEXT(B3,"0")=TEXT(B2,"0"),"OK","NO COINCIDE"))'
    R["A5"], R["B5"] = "Huella del cuaderno", "=CSV!B4"
    R["A6"], R["B6"] = "Sesión", "=CSV!B3"
    R.conditional_formatting.add("B4", FormulaRule(formula=['B4="NO COINCIDE"'], fill=ROJO))
    R.conditional_formatting.add("B4", FormulaRule(formula=['B4="OK"'], fill=VERDE))

    R["A8"] = "Predicho contra obtenido"
    R["A8"].font = NEGRITA
    fila = 9
    for i_csv, clave in res:
        R.cell(fila, 1, ETIQUETAS.get(clave, clave))
        R.cell(fila, 2, f"=CSV!B{i_csv}")
        fila += 1
    R.cell(fila, 1, "Diferencia, obtenido contra predicho (%)")
    R.cell(fila, 2, "=(B11-B10)/B10*100").number_format = "0.0"
    for r in range(9, fila):
        R.cell(r, 2).number_format = "0.0000"
    R.cell(fila + 1, 1, "¿El error obtenido queda dentro de lo predicho?")
    R.cell(fila + 1, 2, '=IF(ABS(B11)<=B10,"Sí","No")')

    fila = fila + 3
    R.cell(fila, 1, "Preguntas del reporte").font = NEGRITA
    preguntas = ["1. ¿Por qué el error obtenido no pasa de medio LSB? Usa tus números.",
                 "2. ¿Qué convertidor (bits) elegirías para registrar la temperatura de un horno de reflujo y por qué?"]
    for p in preguntas:
        fila += 1
        R.cell(fila, 1, p).alignment = Alignment(wrap_text=True, vertical="top")
        R.cell(fila, 2).fill = CAPTURA
        R.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=7)
        R.row_dimensions[fila].height = 60

    # Tabla para la gráfica: los encabezados llevan el nombre del alumno
    R["I1"], R["J1"], R["K1"] = "Lectura", '="Real · "&$B$1', '="Cuantizada · "&$B$1'
    for c in ("I1", "J1", "K1"):
        R[c].font = NEGRITA
    for k in range(len(datos)):
        r_csv = fila_enc + 1 + k
        R.cell(2 + k, 9, k)
        R.cell(2 + k, 10, f"=CSV!B{r_csv}")
        R.cell(2 + k, 11, f"=CSV!C{r_csv}")
    ultimo = 1 + len(datos)

    g = LineChart()
    g.title = "Temperatura real y cuantizada · (tu nombre)"
    g.y_axis.title = "Temperatura (°C)"
    g.x_axis.title = "Lectura (h)"
    g.height, g.width = 9, 22
    g.add_data(Reference(R, min_col=10, max_col=11, min_row=1, max_row=ultimo), titles_from_data=True)
    g.set_categories(Reference(R, min_col=9, min_row=2, max_row=ultimo))
    R.add_chart(g, f"A{fila + 3}")

    R["A" + str(fila + 2)] = ("La gráfica se arma sola al importar tu CSV. Cambia su título por tu nombre "
                              "(doble clic en el título); la leyenda ya lo toma de B1.")
    R["A" + str(fila + 2)].font = Font(italic=True, color="666666")
    wb.save(SALIDA)
    return SALIDA


if __name__ == "__main__":
    print(generar(sys.argv[1]))
